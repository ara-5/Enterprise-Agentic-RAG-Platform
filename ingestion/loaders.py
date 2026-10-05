"""
ingestion/loaders.py
────────────────────
Extracts text from PDF, DOCX, XLSX and PPTX files.

Each file becomes a list of Section(text, page). For PDFs and PPTX the page
is the real page/slide number; for XLSX it is the sheet index; DOCX has no
pagination, so it is returned as a single section with page=1.

Scanned PDF pages (no extractable text layer) are rendered and OCR'd with
Tesseract using the languages in OCR_LANGS (Arabic + English by default).
Requires the Tesseract binary with the `ara` traineddata installed.
"""

from __future__ import annotations

import io
import os
from dataclasses import dataclass
from pathlib import Path
from typing import List

import fitz
import openpyxl
import pytesseract
from docx import Document
from loguru import logger
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".pptx"}

OCR_LANGS    = os.getenv("OCR_LANGS", "ara+eng")
OCR_DPI      = 300
MIN_PAGE_CHARS = 20   # below this a PDF page is treated as scanned


@dataclass(frozen=True)
class Section:
    text: str
    page: int


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_SUFFIXES


def load_document(path: Path) -> List[Section]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".xlsx":
        return _load_xlsx(path)
    if suffix == ".pptx":
        return _load_pptx(path)
    raise ValueError(f"Unsupported file type: {suffix}")


def _load_pdf(path: Path) -> List[Section]:
    reader = PdfReader(str(path))
    sections: List[Section] = []
    ocr_doc = None
    try:
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if len(text.strip()) < MIN_PAGE_CHARS:
                if ocr_doc is None:
                    ocr_doc = fitz.open(str(path))
                try:
                    text = _ocr_pdf_page(ocr_doc[index])
                except pytesseract.TesseractError as e:
                    logger.warning(f"OCR failed on {path.name} page {index + 1}, skipping: {e}")
                    text = ""
            sections.append(Section(text=text, page=index + 1))
    finally:
        if ocr_doc is not None:
            ocr_doc.close()
    return sections


def _ocr_pdf_page(page: fitz.Page) -> str:
    pixmap = page.get_pixmap(dpi=OCR_DPI)
    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
    return ocr_image(image)


def ocr_image(image: Image.Image) -> str:
    try:
        return pytesseract.image_to_string(image, lang=OCR_LANGS)
    except pytesseract.TesseractNotFoundError as e:
        raise RuntimeError(
            "Tesseract is not installed. Install it with the Arabic language pack "
            "(tesseract-ocr and tesseract-ocr-ara) to index scanned PDFs."
        ) from e


def _load_docx(path: Path) -> List[Section]:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [Section(text="\n".join(parts), page=1)]


def _load_xlsx(path: Path) -> List[Section]:
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    sections: List[Section] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            lines = [f"Sheet: {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    lines.append(" | ".join(cells))
            if len(lines) > 1:
                sections.append(Section(text="\n".join(lines), page=index))
    finally:
        workbook.close()
    return sections


def _load_pptx(path: Path) -> List[Section]:
    presentation = Presentation(str(path))
    sections: List[Section] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        if parts:
            sections.append(Section(text="\n".join(parts), page=index))
    return sections
