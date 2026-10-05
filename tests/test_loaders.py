from pathlib import Path

import fitz
import pytesseract
import openpyxl
import pytest
from docx import Document
from pptx import Presentation

from ingestion import loaders
from ingestion.loaders import is_supported, load_document


def _make_pdf(path: Path, pages: list[str | None]) -> None:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text, fontsize=12)
    doc.save(str(path))
    doc.close()


def test_pdf_text_pages_are_extracted_with_page_numbers(tmp_path):
    path = tmp_path / "report.pdf"
    _make_pdf(path, ["First page has enough text to skip OCR.", "Second page also has text here."])

    sections = load_document(path)

    assert [s.page for s in sections] == [1, 2]
    assert "First page" in sections[0].text
    assert "Second page" in sections[1].text


def test_scanned_pdf_pages_are_routed_to_ocr(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    _make_pdf(path, ["Digital text layer on page one, long enough.", None])
    calls = []

    def fake_ocr(page):
        calls.append(page.number)
        return "نص عربي مستخرج"

    monkeypatch.setattr(loaders, "_ocr_pdf_page", fake_ocr)

    sections = load_document(path)

    assert calls == [1]
    assert sections[1].text == "نص عربي مستخرج"
    assert sections[1].page == 2


def test_failing_ocr_page_is_skipped_not_fatal(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    _make_pdf(path, ["Digital text layer on page one, long enough.", None])

    def failing_ocr(page):
        raise loaders.pytesseract.TesseractError(1, "boom")

    monkeypatch.setattr(loaders, "_ocr_pdf_page", failing_ocr)

    sections = load_document(path)

    assert [s.page for s in sections] == [1, 2]
    assert sections[1].text == ""


def test_docx_paragraphs_and_tables_are_extracted(tmp_path):
    path = tmp_path / "memo.docx"
    doc = Document()
    doc.add_paragraph("Quarterly revenue grew by twelve percent.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "Growth"
    doc.save(str(path))

    sections = load_document(path)

    assert len(sections) == 1
    assert "Quarterly revenue grew" in sections[0].text
    assert "Region | Growth" in sections[0].text


def test_xlsx_each_sheet_becomes_a_section(tmp_path):
    path = tmp_path / "budget.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(["Item", "Amount"])
    ws.append(["Rent", 1200])
    other = wb.create_sheet("Empty")
    other.title = "Empty"
    wb.save(str(path))

    sections = load_document(path)

    assert len(sections) == 1
    assert sections[0].page == 1
    assert "Sheet: Expenses" in sections[0].text
    assert "Rent | 1200" in sections[0].text


def test_pptx_slides_and_notes_are_extracted(tmp_path):
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Ship OCR support"
    slide.notes_slide.notes_text_frame.text = "Mention Arabic"
    prs.save(str(path))

    sections = load_document(path)

    assert len(sections) == 1
    assert sections[0].page == 1
    assert "Roadmap" in sections[0].text
    assert "Ship OCR support" in sections[0].text
    assert "Notes: Mention Arabic" in sections[0].text


def test_unsupported_suffix_is_rejected(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("plain text")

    assert not is_supported(path)
    with pytest.raises(ValueError):
        load_document(path)


FIXTURES = Path(__file__).parent / "fixtures"


def _tesseract_has(lang: str) -> bool:
    try:
        return lang in pytesseract.get_languages(config="")
    except pytesseract.TesseractNotFoundError:
        return False


requires_tesseract = pytest.mark.skipif(
    not _tesseract_has("eng"), reason="tesseract binary not installed"
)
requires_arabic = pytest.mark.skipif(
    not _tesseract_has("ara"), reason="tesseract Arabic language data not installed"
)


@requires_tesseract
def test_real_tesseract_ocr_reads_rendered_text(tmp_path, monkeypatch):
    monkeypatch.setattr(loaders, "OCR_LANGS", "eng")
    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "INVOICE NUMBER 4821", fontsize=28)
    doc.save(str(path))
    doc.close()

    text = loaders._ocr_pdf_page(fitz.open(str(path))[0])

    assert "4821" in text


def test_arabic_scan_fixture_has_no_text_layer():
    doc = fitz.open(str(FIXTURES / "arabic_scan.pdf"))
    try:
        assert doc[0].get_text().strip() == ""
    finally:
        doc.close()


@requires_arabic
def test_arabic_scan_is_ocrd_into_arabic_text(monkeypatch):
    monkeypatch.setattr(loaders, "OCR_LANGS", "ara")

    sections = load_document(FIXTURES / "arabic_scan.pdf")

    assert len(sections) == 1
    assert "الإيرادات" in sections[0].text
    assert "2025" in sections[0].text
